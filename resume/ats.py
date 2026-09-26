"""
ATS CHECK — will an applicant-tracking system actually read this file?

    py -m resume.ats "E:/resume/mohith_claude_campous_ambasidor1.pdf"
    py -m resume.ats K_Mohith_Kannan_Resume.pdf --jd job.txt

Why this exists: his best resume looked perfect and contained no text at all — 0 extractable
characters, 0 fonts, 3,273 vector paths per page, produced by `Microsoft: Print To PDF`. Nothing
warned him. Every ATS would have filed it as a blank page. This checker would have said so in its
first line.

What it will not do is print a "match score". The widely quoted "75% of resumes are rejected by
ATS" is traceable to an unsourced sales pitch, and 92% of recruiters surveyed say their ATS does not
auto-reject on formatting or score. An ATS is mostly a database with a search box: what matters is
that the file PARSES, that your skills come out as searchable words, and that the knockout facts
(degree, graduation year, contact) are findable. That is what is checked, and each finding says why.

Per-platform notes come from the open-source ats-screener project's behaviour profiles, which its
own authors call approximations of proprietary systems. They are a map, not a guarantee.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

FAIL, WARN, INFO = "FAIL", "WARN", "INFO"

_STANDARD = {
    "education": r"\beducation\b|\bacademic",
    "skills": r"\bskills?\b|technical skills|core competenc",
    "experience or projects": r"\bexperience\b|\bprojects?\b|\bemployment\b|\bwork history\b|\bselected work\b",
}
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE = re.compile(r"(?:\+?\d{1,3}[\s-]?)?(?:\d[\s-]?){10}")

PLATFORMS = [
    ("Workday", "strict parsing; weighs job-title match heavily — put the target role's words in your headline"),
    ("Taleo", "literal exact keyword match and can auto-reject — use the job's exact terms where true"),
    ("iCIMS", "semantic matching — DOCX is the safest format here"),
    ("Greenhouse", "no auto-scoring by design; the recruiter reads your actual PDF — design matters, text must exist"),
    ("Lever", "stemming, no ranking, early human review — a clean readable page wins"),
    ("SuccessFactors", "normalises skills to a taxonomy — use standard skill names, not your own labels"),
]


@dataclass
class Finding:
    level: str
    what: str
    why: str = ""


@dataclass
class Report:
    path: Path
    kind: str
    pages: int = 0
    chars: int = 0
    text: str = ""
    findings: list[Finding] = field(default_factory=list)

    def add(self, level, what, why=""):
        self.findings.append(Finding(level, what, why))

    @property
    def verdict(self) -> str:
        if any(f.level == FAIL for f in self.findings):
            return FAIL
        return "PASS WITH NOTES" if any(f.level == WARN for f in self.findings) else "PASS"


# ─── PDF ──────────────────────────────────────────────────────────────────────────────────────
def _pdf(path: Path, rep: Report) -> None:
    import fitz  # noqa: PLC0415
    d = fitz.open(str(path))
    rep.pages = len(d)
    pages = [p.get_text() for p in d]
    rep.text = "\n".join(pages)
    rep.chars = len(rep.text.strip())
    producer = (d.metadata or {}).get("producer", "") or ""

    if rep.chars < 200:
        drawn = sum(len(p.get_drawings()) for p in d)
        imgs = sum(len(p.get_images()) for p in d)
        how = (f"every letter was drawn as a shape ({drawn:,} vector paths)" if drawn > 200 else
               f"the pages are images ({imgs} found) — a scan or a screenshot" if imgs else
               "there is no text layer")
        rep.add(FAIL, f"NO READABLE TEXT — {rep.chars} characters could be extracted; {how}.",
                "An ATS reads this as a blank page: no name, no email, no skills. Re-export with your "
                "browser's or Word's own 'Save as PDF', or render it with `py -m resume.render`.")
        if "print to pdf" in producer.lower():
            rep.add(FAIL, f"made with '{producer}'",
                    "Windows' Print-to-PDF driver often outlines text from browsers and design apps. "
                    "Use Chrome/Edge 'Save as PDF' or Word 'Save As → PDF' instead.")
        return

    no_font = [i + 1 for i, p in enumerate(d) if not p.get_fonts()]
    if no_font:
        rep.add(WARN, f"pages with no embedded fonts: {no_font}",
                "Text on those pages may be drawn shapes that a parser cannot read.")

    # Columns: several consecutive text blocks sitting side by side = a multi-column layout. A date
    # pushed right on the SAME line is one line, not a column, so single pairs are ignored.
    side_by_side = 0
    for p in d:
        blocks = [b for b in p.get_text("blocks") if b[4].strip() and b[6] == 0]
        w = p.rect.width
        for a in blocks:
            for b in blocks:
                if a is b or b[0] < a[2]:
                    continue
                overlap = min(a[3], b[3]) - max(a[1], b[1])
                if overlap > 30 and (a[2] - a[0]) > w * 0.2 and (b[2] - b[0]) > w * 0.2:
                    side_by_side += 1
    if side_by_side >= 3:
        rep.add(WARN, "looks like a multi-column layout",
                "Parsers read a PDF in drawing order, and columns come out interleaved. Vendor tests "
                "report two-column PDFs losing far more fields than single-column ones.")

    import contextlib  # noqa: PLC0415
    import io as _io  # noqa: PLC0415
    tables = 0
    for p in d:
        try:
            # PyMuPDF prints an upsell for an optional package on every call; keep the report clean.
            with contextlib.redirect_stdout(_io.StringIO()):
                tables += len(p.find_tables().tables)
        except Exception:  # noqa: BLE001 — find_tables is best-effort
            pass
    if tables:
        rep.add(WARN, f"{tables} table(s) detected", "Many parsers flatten tables out of order.")

    edge = []
    for p in d:
        h = p.rect.height
        for b in p.get_text("blocks"):
            if b[4].strip() and (b[3] < h * 0.035 or b[1] > h * 0.965):
                edge.append(b[4].strip()[:30])
    if edge:
        rep.add(WARN, f"text in the header/footer band: {edge[:2]}",
                "Some ATS skip page headers and footers — keep contact details in the body.")

    if rep.pages > 2:
        rep.add(WARN, f"{rep.pages} pages", "For a student, two pages is the ceiling; recruiters give "
                "the first pass about 7.4 seconds (TheLadders, 2018).")

    title = (d.metadata or {}).get("title", "") or ""
    if re.search(r"[\x00-\x1f]", title):
        rep.add(INFO, f"document title contains a control character: {title!r}",
                "Harmless to parsing, but it shows up in some ATS file lists.")


# ─── DOCX ─────────────────────────────────────────────────────────────────────────────────────
def _docx(path: Path, rep: Report) -> None:
    from docx import Document  # noqa: PLC0415
    doc = Document(str(path))
    body = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            body += [c.text for c in row.cells]
    rep.text = "\n".join(body)
    rep.chars = len(rep.text.strip())
    rep.pages = 0
    if rep.chars < 200:
        rep.add(FAIL, f"NO READABLE TEXT — {rep.chars} characters in the document body",
                "The content may be in text boxes or images, which many parsers skip.")
        return
    if doc.tables:
        rep.add(WARN, f"{len(doc.tables)} table(s)", "Many parsers flatten tables out of order.")
    xml = doc.element.xml
    if "txbxContent" in xml:
        rep.add(WARN, "text boxes found", "Text inside text boxes is skipped by many ATS parsers.")
    hdr = " ".join(p.text for s in doc.sections for p in (s.header.paragraphs + s.footer.paragraphs)).strip()
    if hdr:
        rep.add(WARN, f"text in the page header/footer: {hdr[:40]!r}",
                "Some ATS skip headers and footers — keep contact details in the body.")


# ─── content checks, both formats ────────────────────────────────────────────────────────────
def _content(rep: Report, jd_text: str = "") -> None:
    from .apply import corruption_warnings  # noqa: PLC0415
    from .verify import LEXICON, _present  # noqa: PLC0415

    text = rep.text
    low = re.sub(r"\s+", " ", text.lower())

    emails = _EMAIL.findall(text)
    if not emails:
        rep.add(FAIL, "no email address found", "A recruiter who wants you cannot contact you.")
    if not _PHONE.search(text):
        rep.add(INFO, "no phone number found",
                "Your choice — many Indian recruiters do call first, so consider adding one.")
    for site in ("linkedin", "github"):
        if site not in low:
            rep.add(INFO, f"no {site} link found")

    first = low[:250]
    if emails and emails[0].lower() not in low[:600]:
        rep.add(WARN, "contact details are not near the top", "Parsers look for identity first.")
    if not re.search(r"[a-z]{3,} [a-z]{3,}", first):
        rep.add(WARN, "the first thing extracted is not a name", "Check the reading order.")

    for label, pat in _STANDARD.items():
        if not re.search(pat, low):
            rep.add(WARN, f"no standard '{label}' heading found",
                    "Older parsers classify sections by heading words; use a conventional heading.")
    if re.search(r"\bselected work\b", low) and not re.search(r"\bprojects?\b|\bexperience\b", low):
        rep.add(INFO, "'Selected Work' is a non-standard heading",
                "Most parsers map it to experience, but 'Projects' is the safest label for student work.")

    for w in corruption_warnings(text):
        rep.add(WARN, f"damaged text: {w}", "Keyword search will miss corrupted words.")
    splits = re.findall(r"[A-Za-z]{2,}-\n[a-z]{2,}", text)
    if len(splits) > 3:
        rep.add(INFO, f"{len(splits)} words break across lines at a hyphen",
                "Fine for compound words; a problem only if the hyphen was inserted mid-word.")

    if jd_text:
        jd_low = jd_text.lower()
        asked = [t for t in LEXICON if _present(t, jd_low)]
        have = [t for t in asked if _present(t, low)]
        missing = [t for t in asked if t not in have]
        rep.add(INFO, f"job keywords found in your file: {len(have)}/{len(asked)}"
                + (f" — missing: {', '.join(missing[:10])}" if missing else ""),
                "Exact words are what a recruiter's keyword search and Taleo-style matching find. Add "
                "a missing term only if it is TRUE for you.")


def check(path: str | Path, jd_text: str = "") -> Report:
    path = Path(path)
    kind = path.suffix.lower().lstrip(".")
    rep = Report(path, kind)
    if not path.exists():
        rep.add(FAIL, "file not found")
        return rep
    if kind == "pdf":
        _pdf(path, rep)
    elif kind == "docx":
        _docx(path, rep)
    else:
        rep.add(FAIL, f"'.{kind}' is not a format ATS systems reliably accept",
                "Send a text-based PDF or a DOCX.")
        return rep
    if rep.chars >= 200:
        _content(rep, jd_text)

    name = path.stem
    if re.search(r"\d$", name) or not re.search(r"resume|cv", name, re.I):
        rep.add(INFO, f"file name '{path.name}'",
                "Recruiters see the file name before the page. Something like K_Mohith_Kannan_Resume.pdf.")
    return rep


def format_report(rep: Report) -> str:
    mark = {FAIL: "✗", WARN: "!", INFO: "·"}
    out = [f"\n  {rep.path.name}", f"  {rep.kind.upper()} · "
           + (f"{rep.pages} page(s) · " if rep.pages else "") + f"{rep.chars:,} characters of text",
           f"  VERDICT: {rep.verdict}", ""]
    for lvl in (FAIL, WARN, INFO):
        for f in [x for x in rep.findings if x.level == lvl]:
            out.append(f"  {mark[lvl]} [{lvl}] {f.what}")
            if f.why:
                out.append(f"          {f.why}")
    if rep.verdict != FAIL:
        out += ["", "  How the big platforms treat it (approximations, not guarantees):"]
        out += [f"    {n:<15}{note}" for n, note in PLATFORMS]
    return "\n".join(out) + "\n"


def main() -> int:
    import argparse  # noqa: PLC0415
    ap = argparse.ArgumentParser(description="Will an ATS actually read this resume?")
    ap.add_argument("file", help="a .pdf or .docx")
    ap.add_argument("--jd", default="", help="optional job description file or text, for keyword coverage")
    a = ap.parse_args()
    jd = ""
    if a.jd:
        p = Path(a.jd)
        jd = p.read_text(encoding="utf-8", errors="ignore") if p.exists() else a.jd
    rep = check(a.file, jd)
    print(format_report(rep))
    return 1 if rep.verdict == FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
