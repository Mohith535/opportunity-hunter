"""
Resume Intelligence (Slice 1) — honest resume ↔ job-description analysis.

The whole point is HONESTY. There are three deliberate refusals baked in:

  * No fake "ATS score." Real ATS (Workday/Greenhouse/Lever) don't hand out a number, and a made-up
    "91%" gives false confidence. We report what actually breaks a parser (images, tables, text
    boxes, non-standard headings, multi-column) and a REAL keyword-coverage count ("12 of 18
    must-haves"), never a fabricated percentage.
  * No invented experience. Missing JD keywords are surfaced as QUESTIONS — "do you actually have
    this?" — never auto-added to a resume. Matching a keyword you can't back up is how you fail the
    interview, not pass it.
  * No auto-submit. This analyses and drafts; a human always edits and applies.

Free + light: text extraction via optional pdfminer.six / python-docx (plain .txt needs nothing),
and JD keyword extraction reuses the project's own free LLM chain (filters/llm_scorer.complete).
No paid APIs, no heavy NLP models.

Grounded in how ATS actually parse (2026): exact keyword matches still score highest, so we check
literal coverage; reverse-chronological single-column with standard headings parses best, so those
are what we flag.
"""

from __future__ import annotations

import re
from pathlib import Path

# Standard section headings ATS match against a built-in dictionary. Creative headings get skipped.
_STANDARD_SECTIONS = (
    "experience", "work experience", "employment", "education",
    "skills", "projects", "certifications", "summary", "objective",
)
# Words that are never a useful "must-have keyword" if the model returns them.
_KW_STOP = {
    "experience", "years", "year", "strong", "good", "excellent", "ability", "team",
    "communication", "skills", "knowledge", "understanding", "work", "working", "plus",
    "etc", "the", "and", "with", "for",
}


# ─── 1. text extraction ──────────────────────────────────────────────
def extract_text(path: str | Path) -> str:
    """Plain text from a resume. .txt/.md need nothing; .pdf/.docx use optional libraries with a
    clear install hint if absent (never a cryptic import error)."""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"resume not found: {p}")
    ext = p.suffix.lower()
    if ext in (".txt", ".md"):
        return p.read_text(encoding="utf-8", errors="ignore")
    if ext == ".pdf":
        try:
            from pdfminer.high_level import extract_text as _pdf  # noqa: PLC0415
        except ImportError as e:
            raise RuntimeError("PDF support needs: pip install pdfminer.six") from e
        return _pdf(str(p)) or ""
    if ext == ".docx":
        try:
            import docx  # python-docx  # noqa: PLC0415
        except ImportError as e:
            raise RuntimeError("DOCX support needs: pip install python-docx") from e
        return "\n".join(par.text for par in docx.Document(str(p)).paragraphs)
    raise RuntimeError(f"unsupported resume format '{ext}' — use .pdf, .docx, or .txt")


# ─── 2. honest ATS parse check (what actually breaks parsers) ─────────
def ats_format_check(path: str | Path, text: str) -> list[dict]:
    """Concrete, real parse issues — each with a fix. Severity: critical | high | medium | low."""
    issues: list[dict] = []
    ext = Path(path).suffix.lower()
    words = len(text.split())
    low = text.lower()

    def add(sev, issue, fix):
        issues.append({"severity": sev, "issue": issue, "fix": fix})

    if words < 60:
        add("critical",
            "The parser extracted almost no text — your resume is likely image-based or built from "
            "text boxes an ATS can't read.",
            "Export a real text PDF from your editor (not a scan/screenshot); avoid images/text boxes.")

    if not any(h in low for h in ("experience", "employment")):
        add("high",
            "No standard 'Work Experience' heading found — ATS match headings against a fixed "
            "dictionary and skip creative ones.",
            "Use exact headings: 'Work Experience', 'Education', 'Skills' (not 'Career Journey' etc.).")
    if "skills" not in low:
        add("high",
            "No 'Skills' section — this is exactly where an ATS and a recruiter scan for keyword "
            "matches.",
            "Add a 'Skills' section as a comma-separated list mirroring the job's exact language.")
    if "education" not in low:
        add("medium", "No 'Education' heading detected.", "Add a clearly-labelled 'Education' section.")

    if not re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text):
        add("high", "No email address detected as plain text.",
            "Put your email as plain text near the top — not inside a header, footer, or image.")
    if not re.search(r"\+?\d[\d\s().-]{7,}\d", text):
        add("low", "No phone number detected.", "Add a phone number as plain text.")

    if words > 1100:
        add("medium",
            f"Long resume (~{words} words). ATS and recruiters favour 1 page (2 for senior).",
            "Trim to the most relevant, recent, quantified bullets.")

    if ext == ".docx":
        issues += _docx_structural(path)
    elif ext == ".pdf" and _looks_multicolumn(text):
        add("medium",
            "Possible multi-column layout — Workday's parser scrambles columns and reads content out "
            "of order.",
            "Use a single-column layout; multi-column looks nice to humans but breaks parsers.")
    return issues


def _docx_structural(path: str | Path) -> list[dict]:
    out: list[dict] = []
    try:
        import docx  # noqa: PLC0415
        d = docx.Document(str(path))
        if d.tables:
            out.append({"severity": "high",
                        "issue": f"{len(d.tables)} table(s) found — ATS parse table cells out of "
                                 "document order, scrambling your content.",
                        "fix": "Replace tables with simple single-column bullet lists."})
        if any("image" in r.reltype for r in d.part.rels.values()):
            out.append({"severity": "medium",
                        "issue": "Embedded image(s) found — anything inside an image is invisible to "
                                 "the ATS.",
                        "fix": "Remove images; put everything as selectable text."})
    except Exception:
        pass
    return out


def _looks_multicolumn(text: str) -> bool:
    """Cheap heuristic: many lines with a big internal gap suggest two columns joined per line."""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(lines) < 15:
        return False
    gappy = sum(1 for ln in lines if re.search(r"\S {6,}\S", ln))
    return gappy / len(lines) > 0.35


# ─── 3. JD keyword extraction + honest coverage ──────────────────────
_JD_PROMPT = """From this job description, list the concrete MUST-HAVE hard skills, tools,
technologies, and qualifications an applicant-tracking system would key on — real, matchable terms
(e.g. Python, Docker, REST APIs, machine learning, Kubernetes, SQL). No soft skills, no fluff, no
generic words like "communication" or "team player". Return ONLY a comma-separated list, most
important first, at most 25 items.

JOB DESCRIPTION:
{jd}

Comma-separated must-have keywords:"""


def extract_jd_keywords(jd_text: str) -> list[str]:
    """The JD's matchable must-haves, via the project's free LLM chain. [] if the LLM is unavailable."""
    from filters.llm_scorer import complete  # reuse the free Groq->Cerebras->OpenRouter chain
    out = complete(_JD_PROMPT.format(jd=jd_text[:4000]), max_tokens=300, temperature=0.2)
    if not out:
        return []
    seen, kws = set(), []
    for raw in re.split(r"[,\n;]+", out):
        kw = re.sub(r"^[\s\-*\d.)]+", "", raw).strip().strip('"').strip("`").strip()
        if 2 <= len(kw) <= 40 and kw.lower() not in _KW_STOP and kw.lower() not in seen:
            seen.add(kw.lower())
            kws.append(kw)
    return kws[:25]


def keyword_coverage(resume_text: str, keywords: list[str]) -> tuple[list[str], list[str]]:
    """DETERMINISTIC literal presence check (present, missing) — the honest, verifiable part.
    Exact match is what older ATS reward and what a recruiter's keyword search finds."""
    low = resume_text.lower()
    present, missing = [], []
    for kw in keywords:
        pattern = r"\b" + re.escape(kw.lower().strip()) + r"\b"
        (present if re.search(pattern, low) else missing).append(kw)
    return present, missing


# ─── 4. assemble ─────────────────────────────────────────────────────
def analyze(resume_path: str | Path, jd_text: str) -> dict:
    text = extract_text(resume_path)
    issues = ats_format_check(resume_path, text)
    keywords = extract_jd_keywords(jd_text)
    present, missing = keyword_coverage(text, keywords)
    return {
        "resume_path": str(resume_path),
        "word_count": len(text.split()),
        "issues": issues,
        "keywords": keywords,
        "present": present,
        "missing": missing,
    }


_SEV_ICON = {"critical": "🔴", "high": "🟠", "medium": "🟡", "low": "⚪"}


def format_report(res: dict) -> str:
    lines = ["=" * 60, "RESUME INTELLIGENCE — honest analysis", "=" * 60,
             f"Resume: {res['resume_path']}  ({res['word_count']} words)", ""]

    lines.append("── ATS parse check (what actually breaks parsers) ──")
    if not res["issues"]:
        lines.append("  ✅ No structural parse-killers detected.")
    else:
        order = {"critical": 0, "high": 1, "medium": 2, "low": 3}
        for it in sorted(res["issues"], key=lambda x: order.get(x["severity"], 9)):
            lines.append(f"  {_SEV_ICON.get(it['severity'], '•')} {it['issue']}")
            lines.append(f"       → {it['fix']}")
    lines.append("")

    kws = res["keywords"]
    if not kws:
        lines.append("── JD keyword coverage ──")
        lines.append("  (Could not extract keywords — set an LLM key, or check the JD text.)")
    else:
        cov = len(res["present"])
        lines.append(f"── JD keyword coverage: {cov} / {len(kws)} must-haves present ──")
        lines.append(f"  ✅ Present : {', '.join(res['present']) or '—'}")
        lines.append("")
        lines.append("  ⚠️ Missing — ADD ONLY WHAT YOU'VE GENUINELY DONE (never invent):")
        for kw in res["missing"]:
            lines.append(f"       • {kw}  — do you actually have this? if yes, add it in your own words")
    lines += ["", "This is analysis, not auto-editing. You confirm every claim; you submit. 🧭"]
    return "\n".join(lines)


if __name__ == "__main__":
    # Tiny self-check on inline text (no file, no network needed for the format logic).
    sample = ("John Doe\njohn@example.com\nWork Experience\nBuilt REST APIs in Python.\n"
              "Skills\nPython, FastAPI, Git\nEducation\nB.Tech CSE")
    for it in ats_format_check("x.txt", sample):
        print(it["severity"], "-", it["issue"][:60])
    print("coverage:", keyword_coverage(sample, ["Python", "Docker", "FastAPI"]))
